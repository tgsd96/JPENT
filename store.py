import mysql.connector as mc
import openpyxl as op
import re 

company = 'mar'
#Column configurations for each type of the files, update here to change
configs = {
    'cad' : {
        'start' : '7',
        'get': {
        'billno' : 'B',
        'date' : '',
        'party' : 'C',
        'tinno' : 'D',
        'billvalue' : 'R',
        'sales145' : 'J',
        'vat125' : 'K',
        'sat2' : 'L',
        'sales5' : '',
        'vat4' : '',
        'sat1' : '',
        'roff' : '',
        'disc1' : '',
        'disc2' : '',
        'disc3' : '',
        'disc4' : '',
        'disc5' : '',
        'disc6' : '',
        'disconsale' : '',
        },
        'end' : '0'
    },
    'col' : {
        'start' : '2',
        'get' : {
        'billno' : 'E',
        'date' : '',
        'party' : 'K',
        'tinno' : 'L',
        'billvalue' : 'AB',
        'sales145' : 'AA',
        'vat125' : 'U',
        'sat2' : 'W',
        'sales5' : '',
        'vat4' : '',
        'sat1' : '',
        'roff' : '',
        'disc1' : '',
        'disc2' : '',
        'disc3' : '',
        'disc4' : '',
        'disc5' : '',
        'disc6' : '',
        'disconsale' : '',
        },
        'end' : '0'
    },
    'god' : {
        'start' : '4',
        'get' : {
        'billno' : 'A',
        'date' : '',
        'party' : 'E',
        'tinno' : '',
        'billvalue' : 'W',
        'sales145' : 'S',
        'vat125' : 'T',
        'sat2' : 'P',
        'sales5' : 'Q',
        'vat4' : 'R',
        'sat1' : 'N',
        'roff' : '',
        'disc1' : '',
        'disc2' : '',
        'disc3' : '',
        'disc4' : '',
        'disc5' : '',
        'disc6' : '',
        'disconsale' : '',
        },
        'end' : '0'
    },
    'mar' : {
        'start' : '7',
        'get' : {
        'billno' : 'B',
        'date' : '',
        'party' : 'E',
        'tinno' : 'H',
        'billvalue' : 'AV',
        'sales145' : 'N',
        'vat125' : 'S',
        'sat2' : 'U',
        'sales5' : 'O',
        'vat4' : 'T',
        'sat1' : 'V',
        'roff' : '',
        'disc1' : '',
        'disc2' : '',
        'disc3' : '',
        'disc4' : '',
        'disc5' : '',
        'disc6' : '',
        'disconsale' : '',
        },
        'end' : '0'
    }
}

def dateNormalizer(da):
    # print(da)
    # pattern = re.compile("(\d{1,2})[\/  \-](\d{1,2})[\/ \-](\d{4})")
    # da = pattern.search(da)
    # print(da.group())
    da = da.split('/')
    newdate = da[2]+'-'+da[1]+'-'+da[0]
    return newdate

#load the workbook
print("hi")
wb = op.load_workbook('JPMAR.xlsx')
sheet = wb.get_sheet_by_name(wb.get_sheet_names()[0])
purchases = []

for row in range(int(configs[company]['start']), sheet.max_row+int(configs[company]['end'])):
    purchase = {}
    for key,values in configs[company]['get'].items():
        if(values!=''):
            # print(values+str(row))
            purchase[key] = sheet[values+str(row)].value
        else:
            purchase[key] = 0
    purchases.append(purchase)

print(purchases[0])

# for i in purchases:
    # print(i['party'])

cnx = mc.connect(user="root",password ="root",host="127.0.0.1",database="JPENT")
cur = cnx.cursor()

for i in range(len(purchases)):
    cur.execute("Select custid from mulcust where name = '%s'"%purchases[i]['party'])
    # cur.fetchone()
    # custid = 1
    custid = -1
    for row in cur:
        custid = row[0]
    purchases[i]['custid'] = custid
    # print(custid)
    purchases[i].pop('party')
    purchases[i].pop('tinno')
    purchases[i].pop('date')
    # purchases[i]['date'] = dateNormalizer(str(purchases[i]['date']))
    # purchases[i]['date'] = purchases[i]['date'].date() 
    # print(purchases[i]['date'])
    # purchases[i].pop('date')
    # cur.execute("INSERT INTO PURCHASES()")
    placeholders = ', '.join(["%s"] * len(purchases[i]))
    table = "purchases"
    columns = ', '.join(purchases[i].keys())
    # print(sql)
    sql = ("INSERT INTO %s ( %s ) VALUES ( %s )" % (table, columns, placeholders))
    # print(sql)
    # print(sql)
    cur.execute(sql,list(purchases[i].values()))
    cnx.commit()
    # col_string = ""
    # for key,value in purchases[i]:

# for i in purchases:
#     print(i)